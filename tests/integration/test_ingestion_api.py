"""Integration tests for PDF and DOCX uploads on ``POST /api/v1/ingest/file``."""

from __future__ import annotations

import io
from unittest.mock import AsyncMock

import pytest
from docx import Document
from httpx import AsyncClient


def _minimal_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Report", level=1)
    doc.add_paragraph("Sample DOCX body for ingestion.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_ingest_processes_pdf_file(
    integration_client: AsyncClient,
    mock_ingestion_pipeline: AsyncMock,
) -> None:
    """``POST /api/v1/ingest/file`` must accept ``.pdf`` files and return HTTP 200."""
    mock_ingestion_pipeline.ingest_document.return_value = 5

    response = await integration_client.post(
        "/api/v1/ingest/file",
        files={"file": ("manual.pdf", b"%PDF-1.4 sample", "application/pdf")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert payload["filename"] == "manual.pdf"
    assert payload["chunks_ingested"] == 5


@pytest.mark.asyncio
async def test_ingest_processes_docx_file(
    integration_client: AsyncClient,
    mock_ingestion_pipeline: AsyncMock,
) -> None:
    """``POST /api/v1/ingest/file`` must accept ``.docx`` files and return HTTP 200."""
    mock_ingestion_pipeline.ingest_document.return_value = 4

    response = await integration_client.post(
        "/api/v1/ingest/file",
        files={
            "file": (
                "report.docx",
                _minimal_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert payload["filename"] == "report.docx"
    assert payload["chunks_ingested"] == 4
